"""XLSX -> Markdown，按 Sheet 分节并输出可读表格。"""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path

from .base import ConversionError, ConversionResult, SectionTracker

MAX_ROWS = 20000
MAX_COLUMNS = 200


def convert_xlsx(source: Path, assets_dir: Path) -> ConversionResult:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ConversionError("缺少 openpyxl 依赖，请先执行 ./build.sh") from exc

    try:
        workbook_values = load_workbook(str(source), data_only=True, read_only=False)
        workbook_formulas = load_workbook(str(source), data_only=False, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise ConversionError(f"XLSX 打开失败：{exc}") from exc

    tracker = SectionTracker()
    blocks: list[str] = []
    assets: list[dict] = []
    warnings: list[str] = []
    failures: list[dict] = []

    try:
        for sheet_index, sheet in enumerate(workbook_values.worksheets, start=1):
            heading, section = tracker.heading(2, f"{sheet_index} {sheet.title}", page=sheet_index)
            blocks.append(heading)
            blocks.append("")
            section_id = section["section_id"]

            if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
                warnings.append(
                    f"Sheet「{sheet.title}」过大（{sheet.max_row} 行 × {sheet.max_column} 列），"
                    f"仅转换前 {MAX_ROWS} 行、{MAX_COLUMNS} 列"
                )
            row_limit = min(sheet.max_row or 0, MAX_ROWS)
            column_limit = min(sheet.max_column or 0, MAX_COLUMNS)

            formula_sheet = workbook_formulas[sheet.title]
            merged = _merged_values(sheet, row_limit, column_limit)
            rows: list[list[str]] = []
            for row_index in range(1, row_limit + 1):
                values: list[str] = []
                for column_index in range(1, column_limit + 1):
                    key = (row_index, column_index)
                    if key in merged:
                        values.append(merged[key])
                        continue
                    cell = sheet.cell(row=row_index, column=column_index)
                    value = cell.value
                    if value is None:
                        formula = formula_sheet.cell(row=row_index, column=column_index).value
                        if isinstance(formula, str) and formula.startswith("="):
                            value = formula
                    values.append(_cell_text(value))
                if any(value.strip() for value in values):
                    rows.append(values)

            if rows:
                width = max(len(row) for row in rows)
                rows = [row + [""] * (width - len(row)) for row in rows]
                header = rows[0]
                blocks.append("| " + " | ".join(header) + " |")
                blocks.append("| " + " | ".join("---" for _ in header) + " |")
                for row in rows[1:]:
                    blocks.append("| " + " | ".join(row) + " |")
            else:
                blocks.append("_(空表)_")

            for image_index, image in enumerate(getattr(sheet, "_images", []) or [], start=1):
                reference = _save_sheet_image(image, assets_dir, source, sheet_index, image_index)
                if reference is None:
                    warnings.append(f"Sheet「{sheet.title}」第 {image_index} 张图片无法读取")
                    continue
                assets.append(
                    {
                        "id": f"img_{len(assets) + 1:03d}",
                        "path": reference,
                        "source": f"{sheet.title}/image{image_index}",
                        "section_id": section_id,
                        "page": sheet_index,
                        "markdown_ref": reference,
                    }
                )
                blocks.append("")
                blocks.append(f"![图片]({reference})")
            blocks.append("")
    except Exception as exc:  # noqa: BLE001
        failures.append({"stage": "xlsx", "error": f"{type(exc).__name__}: {exc}"})
        warnings.append(f"部分 Sheet 转换失败：{type(exc).__name__}: {exc}")
    finally:
        workbook_values.close()
        workbook_formulas.close()

    markdown = "\n".join(blocks).strip()
    if not markdown:
        warnings.append("XLSX 未提取到任何 Sheet 内容")

    return ConversionResult(
        markdown=markdown,
        format="xlsx",
        sections=tracker.sections,
        assets=assets,
        warnings=warnings,
        failures=failures,
    )


def _merged_values(sheet, row_limit: int, column_limit: int) -> dict[tuple[int, int], str]:
    result: dict[tuple[int, int], str] = {}
    for merged_range in sheet.merged_cells.ranges:
        min_row, min_col = merged_range.min_row, merged_range.min_col
        max_row, max_col = merged_range.max_row, merged_range.max_col
        if min_row > row_limit or min_col > column_limit:
            continue
        value = _cell_text(sheet.cell(row=min_row, column=min_col).value)
        if not value:
            continue
        for row in range(max(min_row, 1), min(max_row, row_limit) + 1):
            for column in range(max(min_col, 1), min(max_col, column_limit) + 1):
                result[(row, column)] = value
    return result


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.isoformat(sep=" ", timespec="seconds")
    elif isinstance(value, (date, time)):
        text = value.isoformat()
    else:
        text = str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("|", "\\|")
    text = text.replace("\n", "<br />")
    return text.strip()


def _save_sheet_image(image, assets_dir: Path, source: Path, sheet_index: int, image_index: int) -> str | None:
    try:
        data = image._data()
    except Exception:
        return None
    extension = str(getattr(image, "format", "") or "").lower()
    if extension not in {"png", "jpg", "jpeg", "gif", "bmp", "tiff", "wmf", "emf"}:
        extension = "png"
    if extension == "jpeg":
        extension = "jpg"
    filename = f"sheet{sheet_index}_img{image_index}.{extension}"
    path = assets_dir / filename
    path.write_bytes(data)
    return path.relative_to(assets_dir.parent.parent).as_posix()
